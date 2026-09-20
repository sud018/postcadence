"""Read topic names from a file the user provides."""
from __future__ import annotations

import csv
import re
from pathlib import Path

HEADER_WORDS = {"topic", "topics", "topic name", "topic names", "subject", "title"}
MIN_LENGTH = 3
# Strips "1. ", "1) ", "- ", "* ", "• " from the start of a line
BULLET = re.compile(r"^\s*(?:\d+[.)]\s*|[-*•·]\s*)+")
# PDFs sometimes render bullets as "(cid:127)" when the font mapping is missing
CID = re.compile(r"\(cid:\d+\)")


class TopicError(ValueError):
    """The topic file could not be read or held nothing usable."""


def clean(lines: list[str]) -> list[str]:
    """Tidy raw lines into topics: strip bullets, drop junk, drop duplicates."""
    topics: list[str] = []
    seen: set[str] = set()

    for line in lines:
        text = CID.sub("", str(line or ""))
        text = BULLET.sub("", text).strip()
        if len(text) < MIN_LENGTH:
            continue
        if text.lower() in HEADER_WORDS:
            continue
        key = text.lower()
        if key in seen:
            continue
        seen.add(key)
        topics.append(text)

    return topics


def _from_xlsx(path: Path) -> list[str]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    lines = [row[0] for row in sheet.iter_rows(values_only=True) if row and row[0] is not None]
    workbook.close()
    return lines


def _from_docx(path: Path) -> list[str]:
    from docx import Document

    document = Document(str(path))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.extend(cell.text for cell in row.cells)
    return lines


def _from_pdf(path: Path) -> list[str]:
    import pdfplumber

    lines: list[str] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            lines.extend((page.extract_text() or "").splitlines())
    return lines


def _from_csv(path: Path) -> list[str]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return [row[0] for row in csv.reader(f) if row]


def _from_text(path: Path) -> list[str]:
    return path.read_text(encoding="utf-8").splitlines()


READERS = {
    ".xlsx": _from_xlsx,
    ".xlsm": _from_xlsx,
    ".docx": _from_docx,
    ".pdf": _from_pdf,
    ".csv": _from_csv,
    ".txt": _from_text,
    ".md": _from_text,
}


def load_file(path: str | Path) -> list[str]:
    """Read topics from a file. One topic per row, paragraph or line."""
    path = Path(path).expanduser()
    if not path.exists():
        raise TopicError(f"No such file: {path}")

    reader = READERS.get(path.suffix.lower())
    if reader is None:
        raise TopicError(f"Cannot read {path.suffix} files. Supported: {', '.join(sorted(READERS))}")

    try:
        raw = reader(path)
    except TopicError:
        raise
    except Exception as exc:
        raise TopicError(f"Could not read {path.name}: {exc}") from exc

    topics = clean(raw)
    if not topics:
        raise TopicError(f"No topics found in {path.name}.")
    return topics


def load_typed(text: str) -> list[str]:
    """Read topics typed by hand, separated by newlines or semicolons."""
    parts = [piece for line in text.splitlines() for piece in line.split(";")]
    topics = clean(parts)
    if not topics:
        raise TopicError("No topics found in what you typed.")
    return topics
